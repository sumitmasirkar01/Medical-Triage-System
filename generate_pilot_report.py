"""Generate the evidence-backed RSNA Head CT pilot report.

This is the authoritative report generator for the locked 2026-09-30 pilot.
It reads exported Kaggle artifacts and fails when evidence is missing or altered.
Legacy report generators remain available for historical comparison only.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Iterable

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image
from sklearn.metrics import roc_curve


ROOT = Path(__file__).resolve().parent
DEFAULT_ARTIFACT_DIR = ROOT / "artifacts" / "rsna_baseline"
DEFAULT_OUTPUT = ROOT / "docs" / "STAT_Triage_RSNA_Pilot_Report.docx"

# standard_business_brief-inspired token map, resolved explicitly.
PAGE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
BODY_FONT = "Calibri"
BODY_SIZE = 10.5
NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "22313F"
MUTED = "5B6573"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
GREEN = "2E6B4F"
GOLD = "7A5A00"


def resolve_artifact_dir(base: Path) -> Path:
    required = "test_metrics.json"
    if (base / required).exists():
        return base
    candidates = [p.parent for p in base.rglob(required)] if base.exists() else []
    if len(candidates) != 1:
        raise FileNotFoundError(
            f"Expected one {required} beneath {base}; found {len(candidates)}"
        )
    return candidates[0]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_evidence(artifact_dir: Path) -> dict:
    manifest_path = artifact_dir / "bundle_manifest.json"
    bundle = json.loads(manifest_path.read_text(encoding="utf-8"))
    for name, expected in bundle["files"].items():
        path = artifact_dir / name
        if not path.exists():
            raise FileNotFoundError(path)
        if path.stat().st_size != expected["bytes"]:
            raise ValueError(f"Size mismatch for {name}")
        actual_hash = sha256(path)
        if actual_hash != expected["sha256"]:
            raise ValueError(f"SHA-256 mismatch for {name}")

    return {
        "bundle": bundle,
        "metrics": json.loads((artifact_dir / "test_metrics.json").read_text()),
        "operating": json.loads((artifact_dir / "operating_point.json").read_text()),
        "config": json.loads((artifact_dir / "run_config.json").read_text()),
        "history": pd.read_csv(artifact_dir / "training_history.csv"),
        "test": pd.read_csv(artifact_dir / "test_predictions.csv"),
        "calibration": pd.read_csv(artifact_dir / "calibration_predictions.csv"),
        "cohort": pd.read_csv(artifact_dir / "rsna_pilot_manifest_split.csv"),
        "preview": artifact_dir / "rsna_hu_window_preview.png",
        "artifact_dir": artifact_dir,
    }


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run(run, size=None, bold=None, color=None, italic=None, font=BODY_FONT):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = rgb(color)


def set_cell_fill(cell, color: str):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    )


def set_cell_margins(cell, top=90, bottom=90, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is not None:
        tc_pr.remove(tc_mar)
    tc_pr.append(
        parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:start w:w="{start}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:end w:w="{end}" w:type="dxa"/>'
            "</w:tcMar>"
        )
    )


def set_table_geometry(table, widths: Iterable[int]):
    widths = list(widths)
    if sum(widths) != PAGE_WIDTH_DXA:
        raise ValueError(f"Column widths must sum to {PAGE_WIDTH_DXA}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    for name in ("tblW", "tblInd", "tblLayout"):
        old = tbl_pr.find(qn(f"w:{name}"))
        if old is not None:
            tbl_pr.remove(old)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(PAGE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def shade_header(row):
    for cell in row.cells:
        set_cell_fill(cell, NAVY)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_run(run, size=9, bold=True, color=WHITE)


def add_table(doc, headers, rows, widths, alignments=None, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, header in enumerate(headers):
        paragraph = table.cell(0, index).paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(str(header))
        set_run(run, size=9, bold=True, color=WHITE)
    shade_header(table.rows[0])

    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        if row_index % 2:
            for cell in cells:
                set_cell_fill(cell, LIGHT_GRAY)
        for col_index, value in enumerate(values):
            paragraph = cells[col_index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.alignment = (
                alignments[col_index]
                if alignments
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            run = paragraph.add_run(str(value))
            set_run(run, size=font_size, color=INK)

    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_callout(doc, title: str, text: str, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_fill(cell, CALLOUT)
    set_cell_margins(cell, top=130, bottom=130, start=170, end=170)
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="22" w:color="{accent}"/>'
        '<w:top w:val="nil"/><w:bottom w:val="nil"/><w:right w:val="nil"/>'
        "</w:tcBorders>"
    )
    cell._tc.get_or_add_tcPr().append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{title}: ")
    set_run(r, size=10, bold=True, color=accent)
    r = p.add_run(text)
    set_run(r, size=10, color=INK)
    set_table_geometry(table, [PAGE_WIDTH_DXA])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_run(run, size=8.5, color=MUTED)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal.font.size = Pt(BODY_SIZE)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    title = styles["Title"]
    title.font.name = BODY_FONT
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = rgb(NAVY)
    title.paragraph_format.space_after = Pt(5)

    subtitle = styles["Subtitle"]
    subtitle.font.name = BODY_FONT
    subtitle.font.size = Pt(13)
    subtitle.font.color.rgb = rgb(MUTED)
    subtitle.paragraph_format.space_after = Pt(16)

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 12, 5),
        ("Heading 2", 13, BLUE, 10, 4),
        ("Heading 3", 11.5, DARK_BLUE, 8, 3),
    ):
        style = styles[name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("STAT-Triage AI | RSNA pilot evidence report")
    set_run(r, size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("Research prototype | Page ")
    set_run(r, size=8.5, color=MUTED)
    add_page_number(footer)


def build_evaluation_figure(evidence: dict, output: Path):
    test = evidence["test"]
    history = evidence["history"]
    operating = evidence["operating"]
    y_true = test["target_any"].to_numpy()
    y_score = test["score_any"].to_numpy()
    fpr, tpr, _ = roc_curve(y_true, y_score)
    auc = evidence["metrics"]["primary_metrics"]["roc_auc"]

    plt.style.use("seaborn-v0_8-whitegrid")
    fig, axes = plt.subplots(1, 2, figsize=(10.4, 4.1))
    axes[0].plot(fpr, tpr, color="#2E74B5", linewidth=2.3,
                 label=f"Test ROC-AUC = {auc:.3f}")
    axes[0].plot([0, 1], [0, 1], "--", color="#9AA3AD", linewidth=1)
    axes[0].set(xlabel="False-positive rate", ylabel="True-positive rate",
                title="Locked test ROC")
    axes[0].legend(loc="lower right", frameon=True)
    axes[0].set_xlim(0, 1)
    axes[0].set_ylim(0, 1.02)

    axes[1].plot(history["epoch"], history["train_loss"], marker="o",
                 color="#5B6573", label="Training loss")
    axes[1].plot(history["epoch"], history["validation_loss"], marker="o",
                 color="#C46A3B", label="Validation loss")
    selected_epoch = operating["selected_model_epoch"]
    axes[1].axvline(selected_epoch, color="#2E6B4F", linestyle="--",
                    label=f"Selected epoch {selected_epoch}")
    axes[1].set(xlabel="Epoch", ylabel="Weighted BCE loss",
                title="Model selection and overfitting")
    axes[1].legend(loc="upper left", frameon=True, fontsize=8.5)
    fig.tight_layout()
    output.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def build_preview_panels(source: Path, output_dir: Path):
    """Split the tall six-row audit image into two page-safe panels."""
    output_dir.mkdir(parents=True, exist_ok=True)
    image = Image.open(source)
    # The exported six-row audit figure has a natural gutter after row three.
    # Keep that gutter as the split so no scan is cropped across Word pages.
    midpoint = min(1460, image.height - 1)
    panels = []
    for index, box in enumerate(
        ((0, 0, image.width, midpoint),
         (0, midpoint, image.width, image.height)),
        start=1,
    ):
        path = output_dir / f"rsna_hu_window_preview_panel_{index}.png"
        image.crop(box).save(path, optimize=True)
        panels.append(path)
    return panels


def pct(value):
    return f"{100 * value:.1f}%"


def metric_with_ci(metrics, name):
    value = metrics["primary_metrics"][name]
    ci = metrics["patient_bootstrap_95_ci"][name]
    return f"{value:.3f} ({ci['lower']:.3f}-{ci['upper']:.3f})"


def build_report(evidence: dict, output: Path):
    metrics = evidence["metrics"]
    operating = evidence["operating"]
    config = evidence["config"]
    cohort = evidence["cohort"]
    history = evidence["history"]
    artifact_dir = evidence["artifact_dir"]
    figure_path = output.parent / "generated" / "rsna_evaluation_summary.png"
    build_evaluation_figure(evidence, figure_path)
    preview_panels = build_preview_panels(
        evidence["preview"], output.parent / "generated"
    )

    doc = Document()
    configure_document(doc)

    p = doc.add_paragraph("STAT-TRIAGE AI", style="Title")
    p.paragraph_format.space_before = Pt(12)
    p = doc.add_paragraph(
        "Evidence Report: RSNA Intracranial Hemorrhage Slice-Level Pilot",
        style="Subtitle",
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Authors: ")
    set_run(r, bold=True, color=NAVY)
    set_run(p.add_run("Yukti Madan, Barkha Sahu, Sumit Masirkar"), color=INK)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Faculty advisor: ")
    set_run(r, bold=True, color=NAVY)
    set_run(p.add_run("Dr. Monali Gulhane"), color=INK)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("Run date: ")
    set_run(r, bold=True, color=NAVY)
    set_run(p.add_run("30 September 2026 | Kaggle Tesla T4"), color=INK)

    add_callout(
        doc,
        "Evidence status",
        "First locked test evaluation of a 2,400-slice, patient-separated pilot. "
        "This report supports a research prototype; it does not establish 3D study-level "
        "performance, population calibration, clinical efficacy, or regulatory readiness.",
        accent=GOLD,
    )

    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph(
        "A six-output EfficientNet-B0 baseline was trained on three Hounsfield Unit "
        "window channels derived from RSNA 2019 head CT slices. The model checkpoint "
        "was selected using validation data, the operating threshold was fixed on a "
        "separate calibration split, and the test set was evaluated once. The primary "
        "endpoint was slice-level detection of any intracranial hemorrhage."
    )

    primary = metrics["primary_metrics"]
    ci = metrics["patient_bootstrap_95_ci"]
    add_table(
        doc,
        ["Test endpoint", "Estimate", "Patient-bootstrap 95% CI"],
        [
            ("ROC-AUC", f"{primary['roc_auc']:.3f}",
             f"{ci['roc_auc']['lower']:.3f}-{ci['roc_auc']['upper']:.3f}"),
            ("Sensitivity", pct(primary["sensitivity"]),
             f"{pct(ci['sensitivity']['lower'])}-{pct(ci['sensitivity']['upper'])}"),
            ("Specificity", pct(primary["specificity"]),
             f"{pct(ci['specificity']['lower'])}-{pct(ci['specificity']['upper'])}"),
            ("Negative predictive value", pct(primary["npv"]),
             f"{pct(ci['npv']['lower'])}-{pct(ci['npv']['upper'])}"),
        ],
        [3300, 2300, 3760],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
         WD_ALIGN_PARAGRAPH.CENTER],
    )
    doc.add_paragraph(
        "At the locked threshold, the system detected 68 of 77 positive test slices "
        "and missed 9. It also generated 81 false alerts among 165 negative slices. "
        "The result supports continued research but shows that alert burden remains a "
        "central engineering and workflow problem."
    )

    doc.add_heading("1. Study design and evidence chain", level=1)
    split_order = ["train", "validation", "calibration", "test"]
    split_rows = []
    for split in split_order:
        frame = cohort[cohort["split"] == split]
        split_rows.append((
            split.title(), len(frame), frame["patient_id"].nunique(),
            int(frame["any"].sum()), pct(frame["any"].mean()),
        ))
    add_table(
        doc,
        ["Split", "Slices", "Patients", "Any ICH +", "Positive share"],
        split_rows,
        [2200, 1600, 1800, 1800, 1960],
        [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 4,
    )
    doc.add_paragraph(
        "The pilot contains 2,400 sampled slices: 800 hemorrhage-positive and 1,600 "
        "negative. Sampling deliberately increased prevalence before splitting. Patients "
        "were kept within a single split; the exported manifest also shows no study or "
        "series overlap. Because the pilot is slice-sampled, it does not represent full "
        "volumes or the original RSNA prevalence."
    )
    add_callout(
        doc,
        "Locked sequence",
        "Train on 1,682 slices; select the epoch on 232 validation slices; choose the "
        "threshold on 244 calibration slices; evaluate once on 242 test slices from "
        "222 patients. The test result is not used for further tuning.",
    )

    doc.add_heading("2. Input pipeline and model", level=1)
    doc.add_paragraph(
        "DICOM pixels were transformed with RescaleSlope and RescaleIntercept and mapped "
        "to three normalized channels: brain (center 40, width 80), subdural/blood "
        "(center 75, width 215), and bone (center 600, width 2,800). Images were resized "
        "to 224 x 224 and normalized with the ImageNet statistics recorded in the run "
        "configuration. The model produced independent sigmoid scores for epidural, "
        "intraparenchymal, intraventricular, subarachnoid, subdural, and any hemorrhage."
    )
    captions = (
        "Figure 1a. De-identified negative, epidural, and intraparenchymal examples "
        "after physical HU conversion and three-window projection.",
        "Figure 1b. De-identified intraventricular, subarachnoid, and subdural examples. "
        "Slice labels do not constitute lesion localization annotations.",
    )
    for panel, caption in zip(preview_panels, captions):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(panel), width=Inches(5.35))
        p = doc.add_paragraph(caption)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        for run in p.runs:
            set_run(run, size=8.5, italic=True, color=MUTED)

    doc.add_paragraph(
        "The ImageNet-pretrained EfficientNet-B0 classifier was trained for two frozen-"
        "backbone epochs followed by six upper-backbone fine-tuning epochs. Positive "
        "class weighting was capped at 20 for the sparse epidural label. Epoch 4 was "
        "selected with validation any-hemorrhage ROC-AUC 0.8553. Validation loss rose "
        "afterward while training loss continued to decline, indicating overfitting."
    )

    doc.add_heading("3. Locked evaluation", level=1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(figure_path), width=Inches(6.25))
    p = doc.add_paragraph(
        "Figure 2. Locked test ROC curve and training history. The vertical marker shows "
        "the validation-selected checkpoint; later epochs were not retained."
    )
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        set_run(run, size=8.5, italic=True, color=MUTED)

    rows = []
    for name, label in (
        ("ROC-AUC", "roc_auc"),
        ("Average precision", "average_precision"),
        ("Sensitivity", "sensitivity"),
        ("Specificity", "specificity"),
        ("PPV", "ppv"),
        ("NPV", "npv"),
        ("Accuracy", "accuracy"),
    ):
        rows.append((name, metric_with_ci(metrics, label)))
    add_table(
        doc,
        ["Primary endpoint metric", "Estimate (patient-bootstrap 95% CI)"],
        rows,
        [4200, 5160],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )

    add_table(
        doc,
        ["", "Predicted negative", "Predicted positive"],
        [
            ("Actually negative", primary["tn"], primary["fp"]),
            ("Actually positive", primary["fn"], primary["tp"]),
        ],
        [3200, 3080, 3080],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
         WD_ALIGN_PARAGRAPH.CENTER],
    )
    doc.add_paragraph(
        f"The threshold {operating['threshold']:.6f} was chosen on calibration data to "
        f"target at least {pct(operating['target_sensitivity'])} sensitivity. Calibration "
        f"sensitivity/specificity were {pct(operating['calibration_sensitivity'])} and "
        f"{pct(operating['calibration_specificity'])}; test sensitivity/specificity were "
        f"{pct(primary['sensitivity'])} and {pct(primary['specificity'])}. The sigmoid "
        "value is an operating score, not a population-calibrated probability."
    )

    doc.add_page_break()
    doc.add_heading("4. Exploratory subtype results", level=1)
    subtype_rows = []
    for label, result in metrics["exploratory_subtype_metrics"].items():
        subtype_rows.append((
            label.replace("_", " ").title(),
            result["positive_images"],
            f"{result['roc_auc']:.3f}" if result["roc_auc"] is not None else "Undefined",
            result["status"].replace("exploratory; ", ""),
        ))
    add_table(
        doc,
        ["Subtype", "Positive test slices", "ROC-AUC", "Interpretation"],
        subtype_rows,
        [3000, 1900, 1700, 2760],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
         WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    doc.add_paragraph(
        "All subtype analyses are secondary and exploratory. Epidural has only three "
        "positive test slices; intraventricular has 17. Their point estimates are too "
        "unstable for a validated subtype claim. The comparatively strong intraparenchymal "
        "and intraventricular AUCs are hypotheses for a larger prespecified evaluation."
    )

    doc.add_heading("5. Throughput and reproducibility", level=1)
    timing = metrics["timing"]
    doc.add_paragraph(
        f"The 242-slice test cohort completed in {timing['cohort_elapsed_seconds']:.3f} "
        f"seconds at batch size {timing['batch_size']}, an average of "
        f"{1000 * timing['mean_seconds_per_image_at_batch_size']:.2f} ms per image. "
        "This measurement includes batched DICOM decode, preprocessing, device transfer, "
        "and inference. It is a cohort-throughput measurement and must not be presented "
        "as single-study clinical latency or evidence of real-time scanner integration."
    )
    artifact_rows = []
    for name in (
        "test_metrics.json", "test_predictions.csv", "operating_point.json",
        "training_history.csv", "run_config.json", "rsna_pilot_manifest_split.csv",
        "best_efficientnet_b0.pt",
    ):
        entry = evidence["bundle"]["files"][name]
        artifact_rows.append((name, f"{entry['bytes'] / 1024:.1f} KB",
                              entry["sha256"][:16] + "..."))
    add_table(
        doc,
        ["Artifact", "Size", "SHA-256 prefix"],
        artifact_rows,
        [4700, 1500, 3160],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
         WD_ALIGN_PARAGRAPH.LEFT],
        font_size=8.5,
    )
    doc.add_paragraph(
        "The model checkpoint SHA-256 is "
        f"{metrics['model_checkpoint_sha256']}. The report generator verifies every "
        "bundled artifact against bundle_manifest.json before producing this document."
    )

    doc.add_heading("6. Limitations and claim boundaries", level=1)
    limitations = [
        "The endpoint is slice-level any-hemorrhage detection, not complete 3D study-level triage.",
        "The cohort was enriched before splitting; PPV, NPV, accuracy, average precision, and raw scores do not represent source-population calibration.",
        "The test cohort contains 242 slices from 222 patients, producing wide uncertainty intervals and sparse subtype counts.",
        "The checkpoint was initialized from ImageNet EfficientNet-B0 weights. BiomedCLIP, axial-transformer, SPRT, conformal, chest X-ray, and knee MRI claims were not evaluated in this locked run.",
        "The study does not measure radiologist outcomes, clinical workflow benefit, external-site generalization, or regulatory safety/effectiveness.",
        "The locked test result must not be used for additional model or threshold tuning.",
    ]
    for text in limitations:
        p = doc.add_paragraph(text, style="List Bullet")
        p.paragraph_format.space_after = Pt(3)

    add_callout(
        doc,
        "Interpretation",
        "This run replaces prior synthetic 1.0000-AUC demonstrations as the project's "
        "only artifact-backed real-image performance result. The synthetic tri-modal "
        "benchmark remains useful for software demonstrations but is not clinical evidence.",
        accent=GOLD,
    )

    doc.add_heading("7. Next validated milestone", level=1)
    next_steps = [
        "Build complete CT studies from DICOM metadata and preserve all slices per patient and series.",
        "Train and evaluate a study-level baseline on a larger patient-separated cohort with a new untouched test set.",
        "Measure calibration on a cohort sampled to match the intended use population and report reliability diagnostics.",
        "Quantify alert burden and time-to-review in a prespecified queue simulation using the measured operating point.",
        "Evaluate any sequential-stopping design through actual incremental inference before claiming saved compute or latency.",
    ]
    for text in next_steps:
        p = doc.add_paragraph(text, style="List Number")
        p.paragraph_format.space_after = Pt(3)

    doc.add_heading("Sources", level=1)
    sources = [
        "RSNA Intracranial Hemorrhage Detection, 2019 competition data and description: https://www.kaggle.com/c/rsna-intracranial-hemorrhage-detection",
        "DICOM Standard, PS3.5 Data Structures and Encoding: https://dicom.nema.org/medical/dicom/current/output/html/part05.html",
        "Torchvision EfficientNet-B0 model documentation: https://pytorch.org/vision/stable/models/efficientnet.html",
    ]
    for source in sources:
        p = doc.add_paragraph(source, style="List Bullet")
        p.paragraph_format.space_after = Pt(3)

    doc.core_properties.title = "STAT-Triage AI RSNA Pilot Evidence Report"
    doc.core_properties.subject = "Locked slice-level RSNA ICH pilot evaluation"
    doc.core_properties.author = "Yukti Madan; Barkha Sahu; Sumit Masirkar"
    doc.core_properties.keywords = "RSNA, intracranial hemorrhage, CT, research prototype"

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output, figure_path


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifacts", type=Path, default=DEFAULT_ARTIFACT_DIR)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    artifact_dir = resolve_artifact_dir(args.artifacts.resolve())
    evidence = load_evidence(artifact_dir)
    output, figure = build_report(evidence, args.output.resolve())
    print(f"Verified artifacts: {artifact_dir}")
    print(f"Generated report: {output}")
    print(f"Generated figure: {figure}")


if __name__ == "__main__":
    main()
